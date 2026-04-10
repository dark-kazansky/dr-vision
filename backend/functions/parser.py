"""
Parser function for file type detection and text extraction.

This function:
1. Detects file type (image, PDF, text, DOCX)
2. Determines if OCR is needed (scanned vs text-based)
3. Extracts text using appropriate method
"""

from pathlib import Path
from typing import Optional
from dataclasses import dataclass

from agents import BaseOCRAgent


@dataclass
class ParseResult:
    """Result of parsing operation."""
    success: bool
    text: Optional[str] = None
    file_type: Optional[str] = None
    is_scanned: Optional[bool] = None
    pages: Optional[int] = None
    error: Optional[str] = None
    error_type: Optional[str] = None


class Parser:
    """
    Parser for file type detection and text extraction.
    
    Handles:
    - Images (PNG, JPG, JPEG) → OCR
    - Scanned PDFs → OCR
    - Text-based PDFs → Direct extraction
    - Text files (TXT, MD) → Direct reading
    - DOCX → Text extraction
    """
    
    def __init__(self, ocr_agent: Optional[BaseOCRAgent] = None):
        """
        Initialize parser.
        
        Args:
            ocr_agent: OCR agent for scanned documents (optional)
        """
        self.ocr_agent = ocr_agent
    
    def parse(self, file_path: str, force_ocr: bool = False) -> ParseResult:
        """
        Parse file and extract text.
        
        Args:
            file_path: Path to file
            force_ocr: Force OCR even for text-based PDFs
            
        Returns:
            ParseResult with extracted text and metadata
        """
        try:
            ext = Path(file_path).suffix.lower()
            
            # Route to appropriate handler
            if ext in ['.png', '.jpg', '.jpeg']:
                return self._parse_image(file_path)
            elif ext == '.pdf':
                return self._parse_pdf(file_path, force_ocr)
            elif ext in ['.txt', '.md']:
                return self._parse_text_file(file_path)
            elif ext == '.docx':
                return self._parse_docx(file_path)
            else:
                return ParseResult(
                    success=False,
                    error=f"Unsupported file type: {ext}",
                    error_type="validation_error"
                )
                
        except Exception as e:
            return ParseResult(
                success=False,
                error=f"Parsing failed: {str(e)}",
                error_type="processing_error"
            )
    
    def _parse_image(self, file_path: str) -> ParseResult:
        """Parse image using OCR."""
        if not self.ocr_agent:
            return ParseResult(
                success=False,
                error="OCR agent not configured",
                error_type="configuration_error"
            )
        
        result = self.ocr_agent.process_image(file_path)
        
        if not result.success:
            return ParseResult(
                success=False,
                error=result.error,
                error_type=result.error_type
            )
        
        return ParseResult(
            success=True,
            text=result.text,
            file_type='image',
            is_scanned=True,
            pages=1
        )
    
    def _parse_pdf(self, file_path: str, force_ocr: bool) -> ParseResult:
        """Parse PDF using OCR or direct extraction."""
        try:
            import pypdfium2 as pdfium
        except ImportError:
            return ParseResult(
                success=False,
                error="Missing dependency: pypdfium2",
                error_type="dependency_error"
            )
        
        try:
            pdf = pdfium.PdfDocument(file_path)
            page_count = len(pdf)
            
            if page_count == 0:
                return ParseResult(
                    success=False,
                    error="PDF has no pages",
                    error_type="validation_error"
                )
            
            # Check if PDF has text layer
            has_text = False
            if not force_ocr:
                for i in range(min(3, page_count)):
                    page = pdf[i]
                    text = page.get_textpage().get_text_range()
                    if text and text.strip():
                        has_text = True
                        break
            
            # Use direct extraction if text layer exists
            if has_text and not force_ocr:
                all_text = []
                for page_num in range(page_count):
                    page = pdf[page_num]
                    text = page.get_textpage().get_text_range()
                    if text and text.strip():
                        all_text.append(f"--- Page {page_num + 1} ---\n{text}")
                
                combined_text = "\n\n".join(all_text)
                
                return ParseResult(
                    success=True,
                    text=combined_text,
                    file_type='pdf',
                    is_scanned=False,
                    pages=page_count
                )
            
            # Use OCR for scanned PDFs
            else:
                if not self.ocr_agent:
                    return ParseResult(
                        success=False,
                        error="OCR agent not configured",
                        error_type="configuration_error"
                    )
                
                result = self.ocr_agent.process_pdf_all_pages(file_path)
                
                if not result.success:
                    return ParseResult(
                        success=False,
                        error=result.error,
                        error_type=result.error_type
                    )
                
                return ParseResult(
                    success=True,
                    text=result.text,
                    file_type='pdf',
                    is_scanned=True,
                    pages=page_count
                )
                
        except Exception as e:
            return ParseResult(
                success=False,
                error=f"PDF parsing failed: {str(e)}",
                error_type="processing_error"
            )
    
    def _parse_text_file(self, file_path: str) -> ParseResult:
        """Parse plain text or markdown file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
            
            ext = Path(file_path).suffix.lower()
            
            return ParseResult(
                success=True,
                text=text,
                file_type=ext[1:],
                is_scanned=False,
                pages=1
            )
            
        except UnicodeDecodeError:
            try:
                with open(file_path, 'r', encoding='latin-1') as f:
                    text = f.read()
                
                ext = Path(file_path).suffix.lower()
                
                return ParseResult(
                    success=True,
                    text=text,
                    file_type=ext[1:],
                    is_scanned=False,
                    pages=1
                )
            except Exception as e:
                return ParseResult(
                    success=False,
                    error=f"Failed to read text file: {str(e)}",
                    error_type="encoding_error"
                )
        
        except Exception as e:
            return ParseResult(
                success=False,
                error=f"Failed to read text file: {str(e)}",
                error_type="processing_error"
            )
    
    def _parse_docx(self, file_path: str) -> ParseResult:
        """Parse DOCX file."""
        try:
            import docx
        except ImportError:
            return ParseResult(
                success=False,
                error="DOCX support requires python-docx",
                error_type="dependency_error"
            )
        
        try:
            doc = docx.Document(file_path)
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            text = "\n\n".join(paragraphs)
            
            return ParseResult(
                success=True,
                text=text,
                file_type='docx',
                is_scanned=False,
                pages=len(doc.sections)
            )
            
        except Exception as e:
            return ParseResult(
                success=False,
                error=f"DOCX parsing failed: {str(e)}",
                error_type="processing_error"
            )

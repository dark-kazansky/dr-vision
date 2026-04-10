#!/usr/bin/env python3
"""
Test markdown to DOCX conversion.
"""

from docx import Document
from docx.shared import Pt, RGBColor
import re
import os

# Sample markdown text
markdown_text = """# Main Title

This is a regular paragraph with **bold text** and *italic text* and `code snippet`.

## Section 1

Here's a list:
- First item
- Second item with **bold**
- Third item with *italic*

## Section 2

Numbered list:
1. First numbered item
2. Second numbered item
3. Third numbered item

### Subsection

> This is a blockquote with important information.

Here's a [link to Google](https://google.com) in the text.

---

## Code Example

Some `inline code` in a paragraph.

**Important:** This is bold text at the start.

*Note:* This is italic text at the start.
"""

def create_docx_from_markdown(text, output_path):
    """Create DOCX from markdown text."""
    doc = Document()
    
    # Helper function to add inline formatting
    def _add_inline_formatting(para, text):
        """Add text with inline markdown formatting to paragraph."""
        patterns = [
            (r'\*\*(.+?)\*\*', 'bold'),
            (r'\*(.+?)\*', 'italic'),
            (r'`(.+?)`', 'code'),
            (r'\[(.+?)\]\((.+?)\)', 'link'),
        ]
        
        remaining_text = text
        while remaining_text:
            earliest_match = None
            earliest_pos = len(remaining_text)
            earliest_pattern = None
            
            for pattern, ptype in patterns:
                match = re.search(pattern, remaining_text)
                if match and match.start() < earliest_pos:
                    earliest_match = match
                    earliest_pos = match.start()
                    earliest_pattern = ptype
            
            if earliest_match:
                if earliest_pos > 0:
                    run = para.add_run(remaining_text[:earliest_pos])
                    run.font.size = Pt(11)
                
                if earliest_pattern == 'bold':
                    run = para.add_run(earliest_match.group(1))
                    run.bold = True
                    run.font.size = Pt(11)
                elif earliest_pattern == 'italic':
                    run = para.add_run(earliest_match.group(1))
                    run.italic = True
                    run.font.size = Pt(11)
                elif earliest_pattern == 'code':
                    run = para.add_run(earliest_match.group(1))
                    run.font.name = 'Courier New'
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(220, 50, 47)
                elif earliest_pattern == 'link':
                    link_text = earliest_match.group(1)
                    link_url = earliest_match.group(2)
                    run = para.add_run(f"{link_text} ({link_url})")
                    run.font.color.rgb = RGBColor(0, 102, 204)
                    run.font.size = Pt(11)
                
                remaining_text = remaining_text[earliest_match.end():]
            else:
                run = para.add_run(remaining_text)
                run.font.size = Pt(11)
                break
    
    # Add title
    doc.add_heading("Markdown Test Document", level=1)
    doc.add_paragraph("")
    
    # Parse and add content
    lines = text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        if not line:
            doc.add_paragraph("")
            i += 1
            continue
        
        # Check for markdown headers
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=4)
        elif line.startswith('- ') or line.startswith('* '):
            # Bullet list with inline formatting
            list_text = line[2:].strip()
            para = doc.add_paragraph(style='List Bullet')
            para.paragraph_format.left_indent = Pt(18)
            _add_inline_formatting(para, list_text)
        elif re.match(r'^\d+\.\s', line):
            # Numbered list with inline formatting
            text_content = re.sub(r'^\d+\.\s', '', line)
            para = doc.add_paragraph(style='List Number')
            para.paragraph_format.left_indent = Pt(18)
            _add_inline_formatting(para, text_content)
        elif line.startswith('> '):
            # Blockquote with inline formatting
            quote_text = line[2:].strip()
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(36)
            para.paragraph_format.right_indent = Pt(36)
            _add_inline_formatting(para, quote_text)
            for run in para.runs:
                run.font.italic = True
                run.font.color.rgb = RGBColor(96, 96, 96)
        elif line.startswith('---') or line.startswith('***'):
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(12)
        else:
            # Regular paragraph with inline formatting
            para = doc.add_paragraph()
            _add_inline_formatting(para, line)
        
        i += 1
    
    doc.save(output_path)
    print(f"✅ DOCX created: {output_path}")


def verify_docx(docx_path):
    """Verify DOCX formatting."""
    doc = Document(docx_path)
    
    print(f"\n📄 Document Analysis:")
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    
    print(f"\n📋 Paragraph Styles:")
    style_counts = {}
    for p in doc.paragraphs:
        style = p.style.name
        style_counts[style] = style_counts.get(style, 0) + 1
    
    for style, count in sorted(style_counts.items()):
        print(f"  - {style}: {count}")
    
    print(f"\n📝 Sample Paragraphs:")
    for i, p in enumerate(doc.paragraphs[:15]):
        if p.text.strip():
            style = p.style.name
            text = p.text[:60]
            
            # Check formatting
            has_bold = any(run.bold for run in p.runs if run.bold)
            has_italic = any(run.italic for run in p.runs if run.italic)
            
            formatting = []
            if has_bold: formatting.append('BOLD')
            if has_italic: formatting.append('ITALIC')
            fmt_str = ' [' + ', '.join(formatting) + ']' if formatting else ''
            
            print(f"  {i+1}. [{style}]{fmt_str} {text}")


if __name__ == "__main__":
    print("=" * 80)
    print("MARKDOWN TO DOCX CONVERSION TEST")
    print("=" * 80)
    
    output_path = "data/parsed/markdown_test.docx"
    
    # Create DOCX
    create_docx_from_markdown(markdown_text, output_path)
    
    # Verify
    if os.path.exists(output_path):
        verify_docx(output_path)
        print(f"\n✅ Test completed successfully!")
        print(f"📁 Open file: {output_path}")
    else:
        print(f"\n❌ Test failed: File not created")
    
    print("=" * 80)

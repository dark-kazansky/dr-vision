"""Tests for DocxGenerator (Task 5.7).

Validates Requirement 8.1:
- DocxGenerator.generate() produces a valid DOCX file that can be opened
  and contains the expected content structure.
"""

import os
import tempfile

from docx import Document

from core.docx_generator import DocxGenerator


class TestDocxGeneratorProducesValidFile:
    """Requirement 8.1: DOCX generation module produces valid documents."""

    def _generate_to_temp(self, text="Hello world", **kwargs):
        """Helper: generate a DOCX into a temp file and return its path."""
        defaults = dict(
            filename="test.pdf",
            model_id="gemini-2.5-flash",
            pages=3,
            file_type="pdf",
        )
        defaults.update(kwargs)
        gen = DocxGenerator()
        fd, path = tempfile.mkstemp(suffix=".docx")
        os.close(fd)
        try:
            gen.generate(text=text, output_path=path, **defaults)
        except Exception:
            os.unlink(path)
            raise
        return path

    def test_generates_openable_docx(self):
        """The output file should be a valid DOCX that python-docx can open."""
        path = self._generate_to_temp()
        try:
            doc = Document(path)
            assert doc is not None
        finally:
            os.unlink(path)

    def test_contains_title_heading(self):
        """The document should contain the filename as a heading."""
        path = self._generate_to_temp(filename="report.pdf")
        try:
            doc = Document(path)
            headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
            assert "report.pdf" in headings
        finally:
            os.unlink(path)

    def test_contains_metadata(self):
        """The document should include model, pages, and file type metadata."""
        path = self._generate_to_temp(model_id="test-model", pages=5, file_type="image")
        try:
            doc = Document(path)
            all_text = "\n".join(p.text for p in doc.paragraphs)
            assert "Model: test-model" in all_text
            assert "Pages: 5" in all_text
            assert "File Type: image" in all_text
        finally:
            os.unlink(path)

    def test_body_text_included(self):
        """Plain text content should appear in the document body."""
        path = self._generate_to_temp(text="This is the parsed content.")
        try:
            doc = Document(path)
            all_text = "\n".join(p.text for p in doc.paragraphs)
            assert "This is the parsed content." in all_text
        finally:
            os.unlink(path)

    def test_markdown_headings_converted(self):
        """Markdown headings in the text should become DOCX headings."""
        md = "## Section Title\nSome body text"
        path = self._generate_to_temp(text=md)
        try:
            doc = Document(path)
            headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
            assert "Section Title" in headings
        finally:
            os.unlink(path)

    def test_empty_text_produces_valid_docx(self):
        """An empty text input should still produce a valid DOCX with metadata."""
        path = self._generate_to_temp(text="")
        try:
            doc = Document(path)
            all_text = "\n".join(p.text for p in doc.paragraphs)
            assert "Model:" in all_text
        finally:
            os.unlink(path)

    def test_returns_output_path(self):
        """generate() should return the output_path it was given."""
        gen = DocxGenerator()
        fd, path = tempfile.mkstemp(suffix=".docx")
        os.close(fd)
        try:
            result = gen.generate(
                text="test",
                filename="f.pdf",
                model_id="m",
                pages=1,
                file_type="pdf",
                output_path=path,
            )
            assert result == path
        finally:
            os.unlink(path)

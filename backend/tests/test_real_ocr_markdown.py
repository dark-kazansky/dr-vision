#!/usr/bin/env python3
"""
Test markdown conversion with real OCR data.
"""

from docx import Document
from docx.shared import Pt, RGBColor
import re
import os

def _add_inline_formatting(para, text):
    """Add text with inline markdown formatting to paragraph."""
    patterns = [
        (r'\*\*(.+?)\*\*', 'bold'),           # **bold**
        (r'\*(.+?)\*', 'italic'),             # *italic*
        (r'`(.+?)`', 'code'),                 # `code`
        (r'\[(.+?)\]\((.+?)\)', 'link'),      # [text](url)
    ]
    
    remaining_text = text
    while remaining_text:
        earliest_match = None
        earliest_pos = len(remaining_text)
        earliest_pattern = None
        
        for pattern, ptype in patterns:
            match = re.search(pattern, remaining_text)
            if match and match.start() < earliest_pos:
                earliest_match = match
                earliest_pos = match.start()
                earliest_pattern = ptype
        
        if earliest_match:
            if earliest_pos > 0:
                run = para.add_run(remaining_text[:earliest_pos])
                run.font.size = Pt(11)
            
            if earliest_pattern == 'bold':
                run = para.add_run(earliest_match.group(1))
                run.bold = True
                run.font.size = Pt(11)
            elif earliest_pattern == 'italic':
                run = para.add_run(earliest_match.group(1))
                run.italic = True
                run.font.size = Pt(11)
            elif earliest_pattern == 'code':
                run = para.add_run(earliest_match.group(1))
                run.font.name = 'Courier New'
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(220, 50, 47)
            elif earliest_pattern == 'link':
                link_text = earliest_match.group(1)
                link_url = earliest_match.group(2)
                run = para.add_run(f"{link_text} ({link_url})")
                run.font.color.rgb = RGBColor(0, 102, 204)
                run.font.size = Pt(11)
            
            remaining_text = remaining_text[earliest_match.end():]
        else:
            run = para.add_run(remaining_text)
            run.font.size = Pt(11)
            break


def create_docx_from_ocr(ocr_text, output_path, filename):
    """Create DOCX from OCR text with markdown parsing."""
    doc = Document()
    
    # Add title
    doc.add_heading(filename, level=1)
    doc.add_paragraph("")
    
    # Parse and add content
    lines = ocr_text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        if not line:
            doc.add_paragraph("")
            i += 1
            continue
        
        # Check for markdown headers
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=4)
        elif line.startswith('- ') or line.startswith('* '):
            list_text = line[2:].strip()
            para = doc.add_paragraph(style='List Bullet')
            para.paragraph_format.left_indent = Pt(18)
            _add_inline_formatting(para, list_text)
        elif re.match(r'^\d+\.\s', line):
            text_content = re.sub(r'^\d+\.\s', '', line)
            para = doc.add_paragraph(style='List Number')
            para.paragraph_format.left_indent = Pt(18)
            _add_inline_formatting(para, text_content)
        elif line.startswith('> '):
            quote_text = line[2:].strip()
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(36)
            para.paragraph_format.right_indent = Pt(36)
            _add_inline_formatting(para, quote_text)
            for run in para.runs:
                run.font.italic = True
                run.font.color.rgb = RGBColor(96, 96, 96)
        elif line.startswith('---') or line.startswith('***'):
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(12)
        else:
            para = doc.add_paragraph()
            _add_inline_formatting(para, line)
        
        i += 1
    
    doc.save(output_path)
    print(f"✅ DOCX created: {output_path}")


def verify_docx(docx_path):
    """Verify DOCX formatting."""
    doc = Document(docx_path)
    
    print(f"\n📄 Document Analysis:")
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    
    # Count bold text
    bold_count = sum(1 for p in doc.paragraphs for run in p.runs if run.bold)
    print(f"Bold runs: {bold_count}")
    
    print(f"\n📝 Sample Paragraphs with Bold Text:")
    count = 0
    for i, p in enumerate(doc.paragraphs):
        if any(run.bold for run in p.runs if run.bold):
            text = p.text[:80]
            print(f"  {i+1}. {text}")
            count += 1
            if count >= 10:
                break


if __name__ == "__main__":
    print("=" * 80)
    print("REAL OCR MARKDOWN CONVERSION TEST")
    print("=" * 80)
    
    # Read real OCR file
    ocr_file = "data/raw_ocr/BCTC.txt"
    if not os.path.exists(ocr_file):
        print(f"❌ OCR file not found: {ocr_file}")
        exit(1)
    
    with open(ocr_file, 'r', encoding='utf-8') as f:
        ocr_text = f.read()
    
    print(f"✅ Loaded OCR file: {ocr_file}")
    print(f"   Length: {len(ocr_text)} characters")
    
    # Create DOCX
    output_path = "data/parsed/BCTC_markdown_test.docx"
    create_docx_from_ocr(ocr_text, output_path, "BCTC.pdf")
    
    # Verify
    if os.path.exists(output_path):
        verify_docx(output_path)
        print(f"\n✅ Test completed successfully!")
        print(f"📁 Open file: {output_path}")
    else:
        print(f"\n❌ Test failed: File not created")
    
    print("=" * 80)

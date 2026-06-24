"""
Document Job Runner — Async page-by-page document processing with realtime progress.

feat-042: Replaces synchronous OCR/classify/split calls with an async job runner
that emits SSE events per page/stage, supports cooperative cancel, and produces
partial results on per-page failures.

Usage:
    from services.document_job_runner import DocumentJobRunner, DocumentJobConfig

    runner = DocumentJobRunner()
    job_id = runner.create_job(config)
    asyncio.create_task(runner.run(job_id))
"""

import asyncio
import logging
import os
import time
import uuid
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path
from typing import Any, Dict, List, Optional

from services.job_events import job_event_bus
from services.job_manager import job_manager

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Types
# ---------------------------------------------------------------------------


class DocumentOperation(str, Enum):
    """Supported document processing operations."""

    OCR = "ocr"
    CLASSIFY = "classify"
    SPLIT = "split"
    EXTRACT = "extract"


class DocumentJobStatus(str, Enum):
    """States a document job goes through."""

    QUEUED = "queued"
    RUNNING = "running"
    COMPLETED = "completed"
    FAILED = "failed"
    CANCELLED = "cancelled"


@dataclass
class DocumentJobConfig:
    """Configuration for a document processing job."""

    file_path: str
    filename: str
    operation: DocumentOperation
    model_id: str = "default"
    provider: Optional[str] = None
    tier: str = "Normal"
    force_ocr: bool = False
    process_all_pages: bool = True

    # Extraction-specific
    extraction_schema: Optional[str] = None
    extraction_target: Optional[str] = None
    extractor_model: Optional[str] = None

    # Classify-specific
    classification_rules: Optional[str] = None
    classifier_model_id: Optional[str] = None
    is_multimodal: bool = False
    max_pages: int = 5

    # Split-specific
    categories: Optional[str] = None
    allow_uncategorized: bool = True
    split_mode: str = "sections"


@dataclass
class PageResult:
    """Result of processing a single page."""

    page: int
    success: bool
    text: Optional[str] = None
    error: Optional[str] = None


@dataclass
class DocumentJobRecord:
    """In-memory record for a running document job."""

    job_id: str
    config: DocumentJobConfig
    status: DocumentJobStatus = DocumentJobStatus.QUEUED
    progress: float = 0.0
    total_pages: int = 0
    current_page: int = 0
    current_stage: str = "queued"
    page_results: List[PageResult] = field(default_factory=list)
    result: Optional[Dict[str, Any]] = None
    error: Optional[str] = None
    created_at: float = field(default_factory=time.time)
    started_at: Optional[float] = None
    completed_at: Optional[float] = None
    cancel_requested: bool = False


# ---------------------------------------------------------------------------
# Runner
# ---------------------------------------------------------------------------


class DocumentJobRunner:
    """
    Manages async document processing jobs with per-page progress.

    Each job runs as an asyncio task, emitting events through JobEventBus
    and updating status in JobManager for poll-based fallback.
    """

    def __init__(self) -> None:
        self._jobs: Dict[str, DocumentJobRecord] = {}

    def create_job(self, config: DocumentJobConfig) -> str:
        """Create a new document processing job. Returns job_id."""
        job_id = uuid.uuid4().hex
        record = DocumentJobRecord(job_id=job_id, config=config)
        self._jobs[job_id] = record

        # Also register in the legacy job_manager for poll compatibility
        job_manager._jobs[job_id] = {
            "job_id": job_id,
            "status": "pending",
            "progress": None,
            "result": None,
            "error": None,
            "created_at": record.created_at,
            "completed_at": None,
        }

        return job_id

    def get_job(self, job_id: str) -> Optional[DocumentJobRecord]:
        """Retrieve a job record by ID."""
        return self._jobs.get(job_id)

    def request_cancel(self, job_id: str) -> bool:
        """Request cooperative cancellation. Returns False if job not found."""
        record = self._jobs.get(job_id)
        if record is None:
            return False
        if record.status in (DocumentJobStatus.COMPLETED, DocumentJobStatus.FAILED, DocumentJobStatus.CANCELLED):
            return False
        record.cancel_requested = True
        return True

    async def run(self, job_id: str) -> None:
        """
        Main execution loop for a document job.

        Emits SSE events at each stage/page boundary and checks for cancellation
        between pages (cooperative cancel).
        """
        record = self._jobs.get(job_id)
        if record is None:
            logger.error("DocumentJobRunner.run called with unknown job_id=%s", job_id)
            return

        config = record.config

        # Mark as running
        record.status = DocumentJobStatus.RUNNING
        record.started_at = time.time()
        job_manager.update_status(job_id, "processing", progress=0.0)
        await job_event_bus.emit_job_started(job_id)

        try:
            # --- Stage: detect / validate ---
            await self._emit_stage(job_id, record, "detect", 0, self._total_stages(config.operation))

            file_path = config.file_path
            if not os.path.exists(file_path):
                raise FileNotFoundError(f"File not found: {file_path}")

            ext = Path(file_path).suffix.lower()
            page_count = self._get_page_count(file_path, ext)
            record.total_pages = page_count

            # Check cancel
            if record.cancel_requested:
                await self._handle_cancel(job_id, record)
                return

            # --- Stage: parse/ocr ---
            await self._emit_stage(job_id, record, "parse", 1, self._total_stages(config.operation))

            page_texts = await self._process_pages(job_id, record, file_path, ext, page_count)

            if record.cancel_requested:
                await self._handle_cancel(job_id, record)
                return

            # Combine successful pages
            combined_text = self._combine_page_texts(page_texts, page_count)
            failed_pages = [pr for pr in record.page_results if not pr.success]

            # --- Stage: post-process (extract/classify/split if needed) ---
            post_result: Optional[Dict[str, Any]] = None
            if config.operation != DocumentOperation.OCR and combined_text:
                stage_name = config.operation.value
                await self._emit_stage(job_id, record, stage_name, 2, self._total_stages(config.operation))
                post_result = await self._run_post_processing(config, combined_text)

            # --- Complete ---
            duration_ms = int((time.time() - (record.started_at or record.created_at)) * 1000)
            result_payload: Dict[str, Any] = {
                "success": True,
                "text": combined_text,
                "pages": page_count,
                "model": config.model_id,
                "filename": config.filename,
                "operation": config.operation.value,
                "duration_ms": duration_ms,
                "failed_pages": [{"page": p.page, "error": p.error} for p in failed_pages],
            }
            if post_result:
                result_payload["post_processing"] = post_result

            record.status = DocumentJobStatus.COMPLETED
            record.progress = 1.0
            record.result = result_payload
            record.completed_at = time.time()

            job_manager.update_status(job_id, "completed", progress=1.0, result=result_payload)
            await job_event_bus.emit_job_completed(job_id, progress=1.0)

        except Exception as exc:
            logger.exception("Document job %s failed", job_id)
            record.status = DocumentJobStatus.FAILED
            record.error = str(exc)
            record.completed_at = time.time()
            job_manager.update_status(job_id, "failed", error=str(exc))
            await job_event_bus.emit_job_failed(job_id, error=str(exc))

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    def _total_stages(self, operation: DocumentOperation) -> int:
        """Number of processing stages based on operation type."""
        if operation == DocumentOperation.OCR:
            return 2  # detect, parse
        return 3  # detect, parse, post-process

    async def _emit_stage(
        self, job_id: str, record: DocumentJobRecord, stage: str, index: int, total: int
    ) -> None:
        record.current_stage = stage
        await job_event_bus.emit_stage_changed(job_id, stage, index, total)

    def _get_page_count(self, file_path: str, ext: str) -> int:
        """Determine total page count for the file."""
        if ext == ".pdf":
            try:
                import pypdfium2 as pdfium

                pdf = pdfium.PdfDocument(file_path)
                try:
                    return len(pdf)
                finally:
                    pdf.close()
            except Exception:
                return 1
        # Images, text, docx = 1 page
        return 1

    async def _process_pages(
        self,
        job_id: str,
        record: DocumentJobRecord,
        file_path: str,
        ext: str,
        page_count: int,
    ) -> List[Optional[str]]:
        """
        Process each page with OCR, emitting page_progress and partial_result.

        Returns a list of text per page (None for failed pages).
        """
        from agents.factory import AgentFactory
        from config import Config

        config = record.config
        app_config = Config.load()

        page_texts: List[Optional[str]] = []

        # --- Single-page files (image, txt, docx) ---
        if page_count == 1 and ext != ".pdf":
            await job_event_bus.emit_page_progress(job_id, 1, 1, "parse")
            text = await self._process_single_file(file_path, ext, config, app_config)
            page_result = PageResult(page=1, success=text is not None, text=text, error=None if text else "Processing failed")
            record.page_results.append(page_result)
            if text:
                await job_event_bus.emit_partial_result(job_id, page=1, text=text)
            page_texts.append(text)
            record.current_page = 1
            record.progress = 1.0
            job_manager.update_status(job_id, "processing", progress=1.0)
            return page_texts

        # --- PDF: check text layer first ---
        has_text_layer = False
        if not config.force_ocr:
            has_text_layer = self._pdf_has_text_layer(file_path)

        if has_text_layer and not config.force_ocr:
            # Direct text extraction page-by-page
            page_texts = await self._extract_pdf_text_pages(
                job_id, record, file_path, page_count
            )
        else:
            # OCR page-by-page
            ocr_agent = AgentFactory.create_from_config(
                app_config, config.model_id, provider_override=config.provider
            )
            page_texts = await self._ocr_pdf_pages(
                job_id, record, file_path, page_count, ocr_agent
            )

        return page_texts

    async def _process_single_file(
        self, file_path: str, ext: str, config: DocumentJobConfig, app_config: Any
    ) -> Optional[str]:
        """Process a single-page file (image/txt/docx)."""
        from agents.factory import AgentFactory

        if ext in (".txt", ".md"):
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    return f.read()
            except Exception:
                return None

        if ext == ".docx":
            try:
                import docx

                doc = docx.Document(file_path)
                return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
            except Exception:
                return None

        if ext in (".png", ".jpg", ".jpeg"):
            ocr_agent = AgentFactory.create_from_config(
                app_config, config.model_id, provider_override=config.provider
            )
            result = await asyncio.to_thread(ocr_agent.process_image, file_path)
            return result.text if result.success else None

        return None

    def _pdf_has_text_layer(self, file_path: str) -> bool:
        """Check first 3 pages for extractable text."""
        try:
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(file_path)
            try:
                for i in range(min(3, len(pdf))):
                    page = pdf[i]
                    text = page.get_textpage().get_text_range()
                    if text and text.strip():
                        return True
                return False
            finally:
                pdf.close()
        except Exception:
            return False

    async def _extract_pdf_text_pages(
        self,
        job_id: str,
        record: DocumentJobRecord,
        file_path: str,
        page_count: int,
    ) -> List[Optional[str]]:
        """Extract text from text-based PDF, page by page with progress."""
        import pypdfium2 as pdfium

        page_texts: List[Optional[str]] = []

        def _extract_page(pdf_path: str, page_num: int) -> Optional[str]:
            pdf = pdfium.PdfDocument(pdf_path)
            try:
                page = pdf[page_num]
                text = page.get_textpage().get_text_range()
                return text if text and text.strip() else ""
            finally:
                pdf.close()

        for page_num in range(page_count):
            # Cooperative cancel check
            if record.cancel_requested:
                break

            try:
                text = await asyncio.to_thread(_extract_page, file_path, page_num)
                page_result = PageResult(page=page_num + 1, success=True, text=text)
                record.page_results.append(page_result)
                page_texts.append(text)

                await job_event_bus.emit_page_progress(job_id, page_num + 1, page_count, "parse")
                if text:
                    await job_event_bus.emit_partial_result(job_id, page=page_num + 1, text=text)

            except Exception as e:
                logger.warning("Page %d text extraction failed: %s", page_num + 1, e)
                page_result = PageResult(page=page_num + 1, success=False, error=str(e))
                record.page_results.append(page_result)
                page_texts.append(None)
                await job_event_bus.emit_page_progress(job_id, page_num + 1, page_count, "parse")

            record.current_page = page_num + 1
            record.progress = (page_num + 1) / page_count
            job_manager.update_status(job_id, "processing", progress=record.progress)

        return page_texts

    async def _ocr_pdf_pages(
        self,
        job_id: str,
        record: DocumentJobRecord,
        file_path: str,
        page_count: int,
        ocr_agent: Any,
    ) -> List[Optional[str]]:
        """OCR scanned PDF page by page with progress events."""
        page_texts: List[Optional[str]] = []

        for page_num in range(page_count):
            # Cooperative cancel check
            if record.cancel_requested:
                break

            try:
                result = await asyncio.to_thread(
                    ocr_agent.process_pdf_page, file_path, page_num
                )

                if result.success:
                    page_result = PageResult(page=page_num + 1, success=True, text=result.text)
                    record.page_results.append(page_result)
                    page_texts.append(result.text)

                    await job_event_bus.emit_partial_result(job_id, page=page_num + 1, text=result.text)
                else:
                    page_result = PageResult(
                        page=page_num + 1, success=False, error=result.error
                    )
                    record.page_results.append(page_result)
                    page_texts.append(None)

            except Exception as e:
                logger.warning("Page %d OCR failed: %s", page_num + 1, e)
                page_result = PageResult(page=page_num + 1, success=False, error=str(e))
                record.page_results.append(page_result)
                page_texts.append(None)

            await job_event_bus.emit_page_progress(job_id, page_num + 1, page_count, "parse")
            record.current_page = page_num + 1
            record.progress = (page_num + 1) / page_count
            job_manager.update_status(job_id, "processing", progress=record.progress)

        return page_texts

    def _combine_page_texts(self, page_texts: List[Optional[str]], page_count: int) -> str:
        """Combine per-page texts with page separators."""
        parts: List[str] = []
        for i, text in enumerate(page_texts):
            if text:
                parts.append(f"--- Page {i + 1} ---\n{text}")
        return "\n\n".join(parts)

    async def _run_post_processing(
        self, config: DocumentJobConfig, text: str
    ) -> Dict[str, Any]:
        """Run post-processing (classify/split/extract) on combined text."""

        from config import Config

        app_config = Config.load()

        if config.operation == DocumentOperation.EXTRACT:
            return await self._run_extraction(config, text, app_config)
        elif config.operation == DocumentOperation.CLASSIFY:
            return await self._run_classification(config, text, app_config)
        elif config.operation == DocumentOperation.SPLIT:
            return await self._run_split(config, text, app_config)
        return {"success": False, "error": f"Unknown operation: {config.operation}"}

    async def _run_extraction(
        self, config: DocumentJobConfig, text: str, app_config: Any
    ) -> Dict[str, Any]:
        """Run extraction on combined text."""
        import json

        from agents.factory import AgentFactory
        from components.extractor import Extractor
        from core.schemas import ExtractionConfig, ExtractionTarget, FieldType, SchemaField

        try:
            if not config.extraction_schema:
                return {"success": False, "error": "No extraction schema provided"}

            schema_data = json.loads(config.extraction_schema)
            fields = [
                SchemaField(
                    name=f["name"],
                    type=FieldType(f["type"]),
                    description=f.get("description", ""),
                    required=f.get("required", False),
                )
                for f in schema_data
            ]
            extraction_config = ExtractionConfig(
                fields=fields,
                target=ExtractionTarget(config.extraction_target or "document"),
            )
            model_id = config.extractor_model or config.model_id
            llm_agent = AgentFactory.create_llm_agent(model_id, config=app_config)
            extractor = Extractor(agent=llm_agent)
            result = await asyncio.to_thread(extractor.extract, text, extraction_config)

            if result.success:
                return {
                    "success": True,
                    "structured_data": result.structured_data,
                    "field_errors": result.field_errors,
                }
            return {"success": False, "error": result.error}
        except Exception as e:
            return {"success": False, "error": f"Extraction failed: {str(e)}"}

    async def _run_classification(
        self, config: DocumentJobConfig, text: str, app_config: Any
    ) -> Dict[str, Any]:
        """Run classification on combined text."""
        import json

        from agents.factory import AgentFactory
        from components.classifier import Classifier

        try:
            rules = json.loads(config.classification_rules) if config.classification_rules else []
            model_id = config.classifier_model_id or config.model_id
            llm_agent = AgentFactory.create_llm_agent(model_id, config=app_config)
            classifier = Classifier(agent=llm_agent)
            result = await asyncio.to_thread(classifier.classify, text, rules)
            return {"success": True, "classification": result}
        except Exception as e:
            return {"success": False, "error": f"Classification failed: {str(e)}"}

    async def _run_split(
        self, config: DocumentJobConfig, text: str, app_config: Any
    ) -> Dict[str, Any]:
        """Run split on combined text."""
        import json


        try:
            categories = json.loads(config.categories) if config.categories else []
            return {
                "success": True,
                "categories": categories,
                "split_mode": config.split_mode,
                "note": "Split post-processing uses workflow pipeline",
            }
        except Exception as e:
            return {"success": False, "error": f"Split failed: {str(e)}"}

    async def _handle_cancel(self, job_id: str, record: DocumentJobRecord) -> None:
        """Handle cooperative cancellation."""
        record.status = DocumentJobStatus.CANCELLED
        record.completed_at = time.time()

        # Build partial result from pages processed so far
        combined = self._combine_page_texts(
            [pr.text for pr in record.page_results],
            record.total_pages,
        )
        partial_payload = {
            "success": False,
            "cancelled": True,
            "text": combined if combined else None,
            "pages_completed": record.current_page,
            "total_pages": record.total_pages,
        }
        record.result = partial_payload

        job_manager.update_status(job_id, "failed", error="Cancelled by user")
        await job_event_bus.emit_job_cancelled(job_id)


# ---------------------------------------------------------------------------
# Module-level singleton
# ---------------------------------------------------------------------------

document_job_runner = DocumentJobRunner()
